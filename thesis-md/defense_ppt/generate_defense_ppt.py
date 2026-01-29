#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
根据Markdown内容生成答辩PPT的.pptx文件
运行方法: python3 generate_defense_ppt.py
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import re
import os

# 设置中文字体（需要在系统中安装相应字体）
# Windows系统通常使用 'Microsoft YaHei' 或 'SimHei'
# Mac系统通常使用 'PingFang SC' 或 'Heiti TC'
CHINESE_FONT = 'Microsoft YaHei'  # Windows
# CHINESE_FONT = 'PingFang SC'    # Mac

def read_ppt_content(file_path):
    """读取PPT内容的Markdown文件"""
    with open(file_path, 'r', encoding='utf-8') as f:
        content = f.read()
    return content

def parse_slides(content):
    """解析Markdown内容，提取每个幻灯片的内容"""
    slides = []
    # 使用正则表达式分割各个幻灯片
    sections = re.split(r'\n---\n', content)
    
    for section in sections:
        if section.strip():
            lines = section.strip().split('\n')
            slides.append(lines)
    
    return slides

def create_ppt(slides, output_file):
    """创建PPT文件"""
    prs = Presentation()
    
    # 设置幻灯片尺寸为16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    for i, slide_lines in enumerate(slides):
        if not slide_lines:
            continue
        
        # 判断幻灯片类型
        first_line = slide_lines[0].strip()
        
        if first_line.startswith('# Slide 1:') or i == 0:
            # 封面页
            create_title_slide(prs, slide_lines)
        elif first_line.startswith('# Slide 2:') or '目录' in first_line:
            # 目录页
            create_toc_slide(prs, slide_lines)
        elif '结论' in first_line or '致谢' in first_line:
            # 总结页
            create_summary_slide(prs, slide_lines)
        else:
            # 内容页
            create_content_slide(prs, slide_lines)
    
    # 保存文件
    prs.save(output_file)
    print(f"PPT文件已生成: {output_file}")

def create_title_slide(prs, lines):
    """创建封面页"""
    # 空白布局
    layout = prs.slide_layouts[6]  # 空白布局
    slide = prs.slides.add_slide(layout)
    
    # 添加标题
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.333), Inches(1.5))
    title_frame = title_box.text_frame
    title_frame.word_wrap = True
    
    p = title_frame.paragraphs[0]
    p.text = "城市轨道交通变电所高压开关柜"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.name = CHINESE_FONT
    p.alignment = PP_ALIGN.CENTER
    
    p2 = title_frame.add_paragraph()
    p2.text = "在线监测技术的应用与维护策略分析"
    p2.font.size = Pt(36)
    p2.font.bold = True
    p2.font.name = CHINESE_FONT
    p2.alignment = PP_ALIGN.CENTER
    
    # 添加作者信息
    info_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.5), Inches(12.333), Inches(2))
    info_frame = info_box.text_frame
    
    info_text = ["答辩人：况欣睿", "学号：2022010400756", 
                 "专业：高速铁路维修技术", "指导教师：孙乾坤",
                 "完成日期：2025年12月"]
    
    for i, text in enumerate(info_text):
        if i == 0:
            p = info_frame.paragraphs[0]
        else:
            p = info_frame.add_paragraph()
        p.text = text
        p.font.size = Pt(20)
        p.font.name = CHINESE_FONT
        p.alignment = PP_ALIGN.CENTER

def create_toc_slide(prs, lines):
    """创建目录页"""
    layout = prs.slide_layouts[1]  # 标题和内容布局
    slide = prs.slides.add_slide(layout)
    
    # 设置标题
    title = slide.shapes.title
    title.text = "目 录"
    title.text_frame.paragraphs[0].font.size = Pt(40)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.name = CHINESE_FONT
    
    # 设置内容
    body = slide.placeholders[1]
    tf = body.text_frame
    
    toc_items = [
        "1. 研究背景与意义",
        "2. 研究内容与方法",
        "3. 在线监测技术应用",
        "4. 维护策略与故障处理",
        "5. 模拟分析与效果评估",
        "6. 结论与展望"
    ]
    
    tf.clear()  # 清空默认内容
    
    for i, item in enumerate(toc_items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(24)
        p.font.name = CHINESE_FONT
        p.space_before = Pt(18)

def create_content_slide(prs, lines):
    """创建内容页"""
    layout = prs.slide_layouts[1]  # 标题和内容布局
    slide = prs.slides.add_slide(layout)
    
    # 设置标题
    title = slide.shapes.title
    title_text = lines[0].replace('# Slide ', '').replace(':', '').strip()
    title.text = title_text
    title.text_frame.paragraphs[0].font.size = Pt(32)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.name = CHINESE_FONT
    
    # 设置内容
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.clear()
    
    # 跳过标题行，处理剩余内容
    content_lines = [l for l in lines[1:] if l.strip()]
    
    for i, line in enumerate(content_lines):
        line = line.strip()
        if not line:
            continue
            
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        
        # 处理Markdown格式
        line = line.replace('**', '').replace('*', '')
        
        # 处理表格行
        if '|' in line and '故障类型' not in line:
            p.text = line
            p.font.size = Pt(18)
        # 处理普通内容
        elif line.startswith('- '):
            p.text = line
            p.font.size = Pt(20)
        elif line.startswith('1. ') or line.startswith('2. ') or line.startswith('3. '):
            p.text = line
            p.font.size = Pt(20)
        elif line.startswith('**') and '意义' in line:
            p.text = line.replace('**', '')
            p.font.size = Pt(20)
            p.font.bold = True
        elif '相关系数' in line or '经验公式' in line:
            p.text = line
            p.font.size = Pt(18)
            p.font.bold = True
        elif '系统投资' in line or '年均' in line or '投资回收期' in line:
            p.text = line
            p.font.size = Pt(18)
        else:
            p.text = line
            p.font.size = Pt(20)
        
        p.font.name = CHINESE_FONT
        p.space_before = Pt(12)

def create_summary_slide(prs, lines):
    """创建总结页"""
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    
    # 设置标题
    title = slide.shapes.title
    title_text = lines[0].replace('# Slide ', '').replace(':', '').strip()
    title.text = title_text
    title.text_frame.paragraphs[0].font.size = Pt(32)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.name = CHINESE_FONT
    
    # 设置内容
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.clear()
    
    content_lines = [l for l in lines[1:] if l.strip() and not l.startswith('---')]
    
    for i, line in enumerate(content_lines):
        line = line.strip().replace('**', '')
        
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        
        if '感谢' in line:
            p.text = line
            p.font.size = Pt(28)
            p.alignment = PP_ALIGN.CENTER
        elif '创新点' in line:
            p.text = line
            p.font.size = Pt(20)
            p.font.bold = True
        else:
            p.text = line
            p.font.size = Pt(20)
        
        p.font.name = CHINESE_FONT
        p.space_before = Pt(14)

def main():
    """主函数"""
    # 文件路径
    input_file = 'defense_ppt_content.md'
    output_file = 'defense_ppt.pptx'
    
    # 获取当前目录
    current_dir = os.path.dirname(os.path.abspath(__file__))
    input_path = os.path.join(current_dir, input_file)
    output_path = os.path.join(current_dir, output_file)
    
    print("正在读取PPT内容...")
    content = read_ppt_content(input_path)
    
    print("正在解析幻灯片...")
    slides = parse_slides(content)
    print(f"共解析到 {len(slides)} 页幻灯片")
    
    print("正在生成PPT文件...")
    create_ppt(slides, output_path)
    
    print("\nPPT生成完成！")
    print(f"输出文件: {output_path}")

if __name__ == '__main__':
    main()
